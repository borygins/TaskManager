from __future__ import annotations

from datetime import date
from io import BytesIO

from openpyxl import Workbook
from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession
from sqlalchemy.orm import selectinload

from app.models import Client, ClientFunnel
from app.services.task_service import client_to_export_row

EXPORT_HEADERS = [
    "ФИО",
    "Организация",
    "ИНН",
    "Телефон",
    "Комментарий менеджера",
    "Дата звонка",
    "Следующая дата звонка",
    "Дата последней активности менеджера",
    "Статус",
    "Оборот в месяц",
    "MCC",
    "Банк-конкурент",
    "Адрес регистрации",
    "Воронки",
    "Табельный номер",
]
IMPORT_HEADERS = [header for header in EXPORT_HEADERS if header != "Дата последней активности менеджера"]


async def export_clients_to_workbook(session: AsyncSession, manager_id: int | None = None) -> bytes:
    statement = select(Client).options(
        selectinload(Client.funnel_links).joinedload(ClientFunnel.funnel),
        selectinload(Client.assigned_manager),
        selectinload(Client.comments),
    )
    if manager_id:
        statement = statement.where(Client.assigned_manager_id == manager_id)
    statement = statement.order_by(Client.updated_at.desc())
    result = await session.execute(statement)
    clients = list(result.scalars().unique())

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Задачи"
    sheet.append(EXPORT_HEADERS)
    for client in clients:
        row = client_to_export_row(client)
        sheet.append([row.get(header, "") for header in EXPORT_HEADERS])
    for cells in sheet:
        for cell in cells:
            if isinstance(cell.value, str):
                cell.data_type = "s"

    buffer = BytesIO()
    workbook.save(buffer)
    workbook.close()
    return buffer.getvalue()


def build_import_template() -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Шаблон импорта"
    sheet.append(IMPORT_HEADERS)
    sheet.append(
        [
            "Иванов Иван Иванович",
            "ООО Ромашка",
            "7701234567",
            "+79990000000",
            "Первичный контакт",
            date(2026, 3, 25),
            date(2026, 3, 28),
            "Новый",
            "1500000",
            "5411",
            "Банк N",
            "Москва, ул. Примерная, д. 1",
            "Доля рынка, Спарк",
            "01666222",
        ]
    )
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            if isinstance(cell.value, date):
                cell.number_format = "DD.MM.YYYY"
    sheet.freeze_panes = "A2"
    buffer = BytesIO()
    workbook.save(buffer)
    workbook.close()
    return buffer.getvalue()
