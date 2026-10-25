from __future__ import annotations

import asyncio
import logging
import os
from datetime import date, datetime
from io import BytesIO
from pathlib import Path
from tempfile import NamedTemporaryFile

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from sqlalchemy import select, update

from app.core.config import settings
from app.core.sales import POTENTIAL_TURNOVER_LABEL
from app.db.session import SessionLocal
from app.models import ContactReportState, User
from app.services.contacts import ContactRow, aggregate_contacts, list_contacts


REPORT_HEADERS = (
    "ИНН", "Наименование Юр. Лица", POTENTIAL_TURNOVER_LABEL, "Канал поступления",
    "Статус", "Дата активности", "Вид активности", "Дата следующей активности", "Комментарий",
)
logger = logging.getLogger(__name__)
_report_lock = asyncio.Lock()


def build_contact_workbook(rows: list[ContactRow], manager_name: str, period: str) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Журнал активностей"
    sheet.merge_cells("A1:I1")
    sheet["A1"] = f"Журнал активностей — {period}"
    sheet["A1"].font = Font(name="Calibri", size=16, bold=True, color="126B36")
    sheet.row_dimensions[1].height = 26
    sheet.merge_cells("A2:I2")
    sheet["A2"] = f"Менеджер: {manager_name}"
    sheet["A2"].data_type = "s"
    sheet.row_dimensions[2].height = 22
    sheet.append(list(REPORT_HEADERS))
    side = Side(style="thin", color="CCCCCC")
    border = Border(left=side, right=side, top=side, bottom=side)
    widths = (19, 38, 22, 26, 24, 18, 20, 22, 80)
    for column, width in enumerate(widths, start=1):
        cell = sheet.cell(3, column)
        cell.fill = PatternFill("solid", fgColor="008B45")
        cell.font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = border
        sheet.column_dimensions[cell.column_letter].width = width
    sheet.row_dimensions[3].height = 42
    for row_number, row in enumerate(rows, start=4):
        values = (
            row.tax_number, row.company_name, row.potential, row.channel, row.status,
            row.day, row.kind, row.next_day, row.comment,
        )
        for column, value in enumerate(values, start=1):
            cell = sheet.cell(row_number, column, value)
            # Treat all user text literally, including comments starting with '='.
            if isinstance(value, str):
                cell.data_type = "s"
            cell.font = Font(name="Calibri", size=11)
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.border = border
            if row_number % 2 == 0:
                cell.fill = PatternFill("solid", fgColor="F0F6F2")
        sheet.cell(row_number, 1).number_format = "@"
        sheet.cell(row_number, 3).number_format = '[=0]"000000";#,##0.00'
        for column in (6, 8):
            sheet.cell(row_number, column).number_format = "dd.mm.yy"
        lines = max(
            sum(max(1, (len(line) + 69) // 70) for line in row.comment.splitlines()),
            (len(row.company_name) + 31) // 32,
            2,
        )
        sheet.row_dimensions[row_number].height = min(409, lines * 16)
    sheet.freeze_panes = "A4"
    sheet.auto_filter.ref = f"A3:I{max(3, sheet.max_row)}"
    sheet.sheet_view.zoomScale = 85
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.paperSize = sheet.PAPERSIZE_A3
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.print_title_rows = "1:3"
    sheet.print_options.horizontalCentered = True
    buffer = BytesIO()
    workbook.save(buffer)
    workbook.close()
    return buffer.getvalue()


def daily_report_path(manager_id: int, day: date) -> Path:
    return settings.reports_dir / day.strftime("%Y-%m") / f"contacts_{manager_id}_{day.isoformat()}.xlsx"


def write_daily_report(rows: list[ContactRow], manager_name: str, manager_id: int, day: date) -> Path:
    content = build_contact_workbook(rows, manager_name, day.strftime("%d.%m.%Y"))
    destination = daily_report_path(manager_id, day)
    destination.parent.mkdir(parents=True, exist_ok=True)
    temporary_path = None
    try:
        with NamedTemporaryFile(dir=destination.parent, prefix=".contact-report-", suffix=".tmp", delete=False) as output:
            temporary_path = Path(output.name)
            output.write(content)
        os.replace(temporary_path, destination)
    finally:
        if temporary_path and temporary_path.exists():
            temporary_path.unlink()
    return destination


async def generate_daily_report(manager_id: int, day: date) -> bool:
    async with _report_lock:
        async with SessionLocal() as session:
            state = await session.get(ContactReportState, (manager_id, day))
            if state is None:
                return True
            revision = state.revision
            manager = await session.get(User, manager_id)
            rows = aggregate_contacts(await list_contacts(session, manager_id, day, day))
            manager_name = manager.full_name
        try:
            await asyncio.to_thread(write_daily_report, rows, manager_name, manager_id, day)
        except Exception:
            logger.exception("Cannot generate contact report for manager %s, day %s", manager_id, day)
            async with SessionLocal() as session:
                await session.execute(update(ContactReportState).where(
                    ContactReportState.manager_id == manager_id,
                    ContactReportState.activity_date == day,
                ).values(pending=True, last_error="Не удалось обновить Excel. Закройте файл и проверьте доступ к папке reports."))
                await session.commit()
            return False
        async with SessionLocal() as session:
            await session.execute(update(ContactReportState).where(
                ContactReportState.manager_id == manager_id,
                ContactReportState.activity_date == day,
                ContactReportState.revision == revision,
            ).values(pending=False, last_error="", generated_at=datetime.now()))
            await session.commit()
        return True


async def retry_pending_reports() -> None:
    async with SessionLocal() as session:
        pending = (await session.execute(select(
            ContactReportState.manager_id, ContactReportState.activity_date,
        ).where(ContactReportState.pending.is_(True)))).all()
    for manager_id, day in pending:
        await generate_daily_report(manager_id, day)
