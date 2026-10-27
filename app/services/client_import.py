from __future__ import annotations

import hashlib
import json
import re
from collections import Counter
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from io import BytesIO
from zipfile import ZipFile

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from sqlalchemy import delete, select
from sqlalchemy.ext.asyncio import AsyncSession
from sqlalchemy.orm import selectinload

from app.core.sales import CLIENT_STATUSES, CLIENT_STATUS_SET, POTENTIAL_TURNOVER_LABEL
from app.core.security import normalize_employee_id
from app.models import Client, ClientFunnel, Comment, Funnel, User
from app.services.task_service import append_daily_activity

MAX_FILE_BYTES = 5 * 1024 * 1024
MAX_ROWS = 5000
IMPORT_POLICY_VERSION = 8
FIELDS = {
    "ЛПР": "full_name", "Организация": "company_name", "Телефон": "phone",
    "Дата звонка": "last_call_date", "Следующая дата звонка": "next_call_date",
    "Статус": "status", POTENTIAL_TURNOVER_LABEL: "monthly_turnover", "MCC": "mcc_code",
    "Банк-конкурент": "competitor_bank", "Адрес регистрации": "registration_address",
}
HEADER_ALIASES = {"ФИО": "ЛПР", "Оборот в месяц": POTENTIAL_TURNOVER_LABEL}
DATES = {"last_call_date", "next_call_date"}
PROTECTED = DATES | {"status", "phone", "full_name", "company_name"}
LABELS = {value: key for key, value in FIELDS.items()} | {
    "assigned_manager_id": "Ответственный (табельный)", "funnels": "Воронки", "comment": "Комментарий",
}
LIMITS = {"full_name": 255, "company_name": 255, "phone": 64, "status": 64,
          "mcc_code": 32, "competitor_bank": 255, "registration_address": 500}


def normalized(value: str) -> str:
    return " ".join(value.split())


def scalar_text(value) -> str:
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.isoformat(sep=" ")
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def date_text(value) -> str:
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    raw = scalar_text(value)
    for fmt in ("%Y-%m-%d", "%d.%m.%Y"):
        try:
            return datetime.strptime(raw, fmt).date().isoformat()
        except ValueError:
            pass
    raise ValueError("ожидается дата Excel, ДД.ММ.ГГГГ или ГГГГ-ММ-ДД")


def phone_key(value: str) -> str:
    digits = re.sub(r"\D", "", value)
    if len(digits) == 11 and digits[0] in "78":
        return "7" + digits[1:]
    return normalized(value)


def fingerprint(value) -> str:
    return hashlib.sha256(json.dumps(value, sort_keys=True, ensure_ascii=False).encode()).hexdigest()


def parse_workbook(content: bytes) -> dict:
    if len(content) > MAX_FILE_BYTES:
        raise ValueError("Размер файла не должен превышать 5 МБ.")
    try:
        with ZipFile(BytesIO(content)) as archive:
            if sum(info.file_size for info in archive.infolist()) > 32 * 1024 * 1024:
                raise ValueError("Распакованный Excel слишком большой (более 32 МБ).")
        workbook = load_workbook(BytesIO(content), read_only=True, data_only=False)
    except Exception as exc:
        raise ValueError(f"Не удалось прочитать XLSX: {exc}") from exc
    try:
        sheet = workbook.active
        if sheet.max_column and sheet.max_column > 64:
            raise ValueError("Допускается не более 64 колонок.")
        iterator = sheet.iter_rows()
        header_cells = next(iterator, ())
        headers = [scalar_text(cell.value) for cell in header_cells]
        canonical_headers = [HEADER_ALIASES.get(name, name) for name in headers]
        nonempty = [name for name in canonical_headers if name]
        if "ИНН" not in canonical_headers or len(nonempty) != len(set(nonempty)):
            raise ValueError("Нужна колонка «ИНН»; названия колонок не должны повторяться.")
        rows = []
        for number, cells in enumerate(iterator, start=2):
            if number > MAX_ROWS + 1:
                raise ValueError(f"Допускается не более {MAX_ROWS} строк. Разделите файл.")
            raw = [date_text(cell.value) if isinstance(cell.value, (date, datetime)) else scalar_text(cell.value) for cell in cells]
            if not any(raw):
                continue
            values = {name: cells[index].value for index, name in enumerate(canonical_headers) if name}
            row = {"row": number, "raw": raw, "inn": scalar_text(values.get("ИНН")), "values": {}, "error": ""}
            try:
                if any(cell.data_type in {"f", "e"} for cell in cells):
                    raise ValueError("Формулы и ошибки Excel не допускаются; вставьте значения.")
                if not re.fullmatch(r"(?:[0-9]{10}|[0-9]{12})", row["inn"]):
                    raise ValueError("ИНН должен содержать 10 или 12 цифр. Храните ИНН как текст.")
                for column, field in FIELDS.items():
                    value = values.get(column)
                    raw_value = scalar_text(value)
                    if not raw_value:
                        continue
                    if field in DATES:
                        try:
                            raw_value = date_text(value)
                        except ValueError as exc:
                            raise ValueError(f"{column}: {exc}") from exc
                    elif field == "monthly_turnover":
                        try:
                            amount = Decimal(raw_value.replace(" ", "").replace("\u00a0", "").replace(",", "."))
                        except InvalidOperation as exc:
                            raise ValueError(f"{POTENTIAL_TURNOVER_LABEL}: некорректное число.") from exc
                        if not amount.is_finite() or not 0 <= amount <= Decimal("999999999999.99") or amount != amount.quantize(Decimal("0.01")):
                            raise ValueError(f"{POTENTIAL_TURNOVER_LABEL}: неотрицательное число до 999999999999.99, не более двух знаков после запятой.")
                        raw_value = format(amount, ".2f")
                    elif len(raw_value) > LIMITS.get(field, 10000):
                        raise ValueError(f"{column}: слишком длинное значение.")
                    row["values"][field] = raw_value
                employee = scalar_text(values.get("Табельный номер"))
                if employee:
                    row["values"]["employee_id"] = normalize_employee_id(employee)
                comment = scalar_text(values.get("Комментарий менеджера"))
                if len(comment) > 10000:
                    raise ValueError("Комментарий: не более 10000 символов.")
                if comment:
                    row["values"]["comment"] = comment
                names = scalar_text(values.get("Воронки"))
                if names:
                    names = sorted({normalized(name) for name in names.split(",") if normalized(name)})
                    if not names or any(len(name) > 255 for name in names):
                        raise ValueError("Укажите названия воронок через запятую (до 255 символов каждое).")
                    row["values"]["funnels"] = names
            except ValueError as exc:
                row["error"] = str(exc)
            rows.append(row)
        if not rows:
            raise ValueError("В файле нет строк клиентов.")
        # Reject conflicting repetitions as a group, never pick an arbitrary winner.
        groups = {}
        for row in rows:
            groups.setdefault(row["inn"], []).append(row)
        for group in groups.values():
            if len(group) <= 1:
                continue
            if any(row["error"] for row in group) or len({fingerprint(row["values"]) for row in group}) > 1:
                for row in group:
                    row["error"] = "Повтор ИНН с различающимися или ошибочными данными в строках: " + ", ".join(str(r["row"]) for r in group)
            else:
                for row in group[1:]:
                    row["duplicate"] = group[0]["row"]
        return {"headers": canonical_headers, "rows": rows}
    except ValueError:
        raise
    except Exception as exc:
        raise ValueError("Не удалось прочитать строки Excel. Проверьте целостность файла и сохраните его как XLSX.") from exc
    finally:
        workbook.close()


async def build_plan(session: AsyncSession, source: dict) -> dict:
    inns = list({row["inn"] for row in source["rows"] if not row["error"]})
    clients = {}
    for start in range(0, len(inns), 400):
        items = await session.scalars(select(Client).where(Client.tax_number.in_(inns[start:start + 400])).options(
            selectinload(Client.comments), selectinload(Client.funnel_links).selectinload(ClientFunnel.funnel),
        ))
        clients.update({client.tax_number: client for client in items})
    users = list(await session.scalars(select(User)))
    by_employee = {user.employee_id: user for user in users}
    by_id = {user.id: user for user in users}
    funnels = {f.name: f.id for f in await session.scalars(select(Funnel))}
    funnel_names = {}
    for name in funnels:
        funnel_names.setdefault(normalized(name).casefold(), []).append(name)
    new_funnels = {}
    entries = []
    for source_row in source["rows"]:
        row = {"row": source_row["row"], "inn": source_row["inn"], "action": "error",
               "error": source_row["error"], "changes": [], "warnings": [], "notes": [], "updates": {}}
        entries.append(row)
        if row["error"]:
            continue
        if source_row.get("duplicate"):
            row.update(action="duplicate", error=f"Повтор строки {source_row['duplicate']}: одинаковые данные.")
            continue
        incoming = dict(source_row["values"])
        if "funnels" in incoming:
            resolved = set()
            for name in incoming["funnels"]:
                key = normalized(name).casefold()
                matches = funnel_names.get(key, [])
                if name in funnels:
                    resolved.add(name)
                elif len(matches) == 1:
                    resolved.add(matches[0])
                elif len(matches) > 1:
                    row["error"] = f"Название воронки «{name}» неоднозначно. Укажите точное название из базы."
                    break
                else:
                    resolved.add(new_funnels.setdefault(key, name))
            if row["error"]:
                continue
            incoming["funnels"] = sorted(resolved)
        client = clients.get(row["inn"])
        manager = by_employee.get(incoming.get("employee_id"))
        if "employee_id" in incoming and (manager is None or manager.role != "manager" or not manager.is_active):
            row["error"] = "Табельный номер не соответствует действующему менеджеру."
            continue
        if not client and not (incoming.get("company_name") or incoming.get("full_name")):
            row["error"] = "Для нового клиента укажите организацию или ЛПР."
            continue
        if "status" in incoming and incoming["status"] not in CLIENT_STATUS_SET:
            row["error"] = "Неизвестный статус клиента."
            continue
        current = {field: scalar_text(getattr(client, field)) for field in FIELDS.values()} if client else {}
        if client and client.monthly_turnover is not None:
            current["monthly_turnover"] = format(client.monthly_turnover, ".2f")
        row["snapshot"] = fingerprint({
            "id": client.id if client else None, "fields": current,
            "updated_at": scalar_text(client.updated_at) if client else None,
            "manager": client.assigned_manager_id if client else None,
            "funnels": sorted((link.funnel_id, link.funnel.name) for link in client.funnel_links) if client else [],
            "comments": sorted((c.id, c.body) for c in client.comments) if client else [],
        })
        if not client:
            incoming.setdefault("full_name", "")
            incoming.setdefault("status", CLIENT_STATUSES[0])
        for field in FIELDS.values():
            if field not in incoming:
                continue
            old, new = current.get(field, ""), incoming[field]
            compare = phone_key if field == "phone" else normalized
            if compare(old) == compare(new):
                continue
            if client and field in PROTECTED and old:
                row["warnings"].append(f"{LABELS[field]}: в базе «{old}», в файле «{new}». Сохранено значение из базы: заполненное поле защищено от замены при импорте.")
                continue
            row["updates"][field] = new
            row["changes"].append({"label": LABELS[field], "old": old, "new": new,
                                   "explanation": "Заполнение пустого поля." if not old else "Замена прежнего значения значением из файла."})
        if manager:
            old_manager = by_id.get(client.assigned_manager_id) if client else None
            # Include the target identity in the plan to detect role/assignment changes before apply.
            row["target_manager"] = [manager.id, manager.employee_id, manager.role, manager.is_active]
            if not client or client.assigned_manager_id != manager.id:
                row["updates"]["assigned_manager_id"] = manager.id
                row["changes"].append({"label": LABELS["assigned_manager_id"], "old": old_manager.employee_id if old_manager else "", "new": manager.employee_id,
                                       "explanation": "Назначение ответственного по табельному номеру из файла."})
        if "funnels" in incoming:
            names = incoming["funnels"]
            row["target_funnels"] = [[name, funnels.get(name)] for name in names]
            old_names = sorted(link.funnel.name for link in client.funnel_links) if client else []
            if old_names != names:
                row["updates"]["funnels"] = names
                row["changes"].append({"label": "Воронки (полная замена)", "old": ", ".join(old_names), "new": ", ".join(names),
                                       "explanation": "Прежние связи заменяются: останутся только перечисленные в файле воронки."})
        comment = incoming.get("comment")
        if comment and (not client or normalized(comment) not in {normalized(c.body) for c in client.comments}):
            row["updates"]["comment"] = comment
            row["changes"].append({"label": "Новый комментарий в историю", "old": "", "new": comment,
                                   "explanation": "Добавление отдельной записи. Прежние комментарии сохраняются, звонок не создаётся."})
        elif comment:
            row["notes"].append("Комментарий уже есть в истории: повторно не добавляется.")
        row["action"] = "create" if not client else "update" if row["updates"] else "unchanged"
        row["client_id"] = client.id if client else None
    counts = Counter(row["action"] for row in entries)
    return {"version": IMPORT_POLICY_VERSION, "entries": entries, "total": len(entries),
            "counts": {key: counts[key] for key in ("create", "update", "unchanged", "error", "duplicate")},
            "apply_count": counts["create"] + counts["update"]}


async def apply_plan(session: AsyncSession, actor: User, plan: dict) -> None:
    if actor.role != "boss":
        raise PermissionError("Импорт доступен только руководителю.")
    for row in plan["entries"]:
        if row["action"] not in {"create", "update"}:
            continue
        updates = dict(row["updates"])
        names, comment = updates.pop("funnels", None), updates.pop("comment", None)
        for field in DATES & updates.keys():
            updates[field] = date.fromisoformat(updates[field])
        if "monthly_turnover" in updates:
            updates["monthly_turnover"] = Decimal(updates["monthly_turnover"])
        if row["action"] == "create":
            client = Client(tax_number=row["inn"], **updates)
            session.add(client)
        else:
            client = await session.get(Client, row["client_id"])
            for field, value in updates.items():
                setattr(client, field, value)
        await session.flush()
        if names is not None:
            await session.execute(delete(ClientFunnel).where(ClientFunnel.client_id == client.id))
            for name in names:
                funnel = await session.scalar(select(Funnel).where(Funnel.name == name))
                if funnel is None:
                    funnel = Funnel(name=name)
                    session.add(funnel)
                    await session.flush()
                session.add(ClientFunnel(client_id=client.id, funnel_id=funnel.id))
        if comment:
            session.add(Comment(client_id=client.id, author_id=actor.id, body=comment, source="import"))
        await append_daily_activity(client.id, actor, [
            "Импорт Excel: " + change["label"] + ": " + str(change["old"] or "—") + " → " + str(change["new"])
            for change in row["changes"]
        ], session)


def report_workbook(source: dict, plan: dict, state: str) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Отчёт импорта"
    report_headers = ["Строка исходного файла", "Результат импорта", "Изменения импорта",
                      "Пояснение импорта", "Ошибка импорта", "Состояние загрузки"]
    diagnostic = set(report_headers)
    indices = [index for index, name in enumerate(source["headers"]) if name not in diagnostic]
    headers = [source["headers"][index] for index in indices] + report_headers
    sheet.append(headers)
    state_label = {"pending": "Предварительная проверка — изменения ещё не применены",
                   "applied": "Импорт завершён", "cancelled": "Отменено — изменения не применены",
                   "stale": "Проверка устарела — изменения не применены"}[state]
    results = {"unchanged": "Без изменений", "error": "Ошибка", "duplicate": "Повтор — пропущен"}
    if state == "applied":
        results.update(create="Создан", update="Обновлён")
    elif state == "pending":
        results.update(create="Будет создан", update="Будет обновлён")
    else:
        results.update(create="Создание не применено", update="Обновление не применено")
    entries = {entry["row"]: entry for entry in plan["entries"]}
    for row in source["rows"]:
        entry = entries[row["row"]]
        changes = "\n".join(
            f"{change['label']}: {change['old'] or '—'} → {change['new']}"
            for change in entry["changes"]
        )
        notes = entry["warnings"] + entry["notes"]
        if entry["action"] == "duplicate":
            notes = notes + [entry["error"]]
        sheet.append([row["raw"][index] for index in indices] + [
            row["row"], results[entry["action"]], changes, "\n".join(notes),
            entry["error"] if entry["action"] == "error" else "", state_label,
        ])
    for cells in sheet:
        for cell in cells:
            if isinstance(cell.value, str):
                cell.data_type = "s"
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="166534")
    for index, header in enumerate(headers, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = (
            60 if header in {"Изменения импорта", "Пояснение импорта", "Ошибка импорта", "Комментарий менеджера"}
            else 38 if header == "Состояние загрузки" else 24
        )
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    buffer = BytesIO()
    workbook.save(buffer)
    workbook.close()
    return buffer.getvalue()
