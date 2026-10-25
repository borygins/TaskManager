from __future__ import annotations

import asyncio
import html
import json
import re
import unittest
from datetime import date, datetime, timedelta
from io import BytesIO
from unittest.mock import patch

from openpyxl import Workbook, load_workbook
from sqlalchemy import func, select

import test_contacts
from app.api.deps import get_current_user
from app.main import app
from app.models import Client, ClientFunnel, Comment, ContactActivity, Funnel, User
from app.models.entities import ClientImport
from app.services.client_import import MAX_FILE_BYTES, apply_plan, parse_workbook
from app.services.excel import build_import_template, export_clients_to_workbook


class ImportTests(unittest.IsolatedAsyncioTestCase):
    asyncTearDown = test_contacts.ContactWorkflowTests.asyncTearDown
    request = test_contacts.ContactWorkflowTests.request

    async def asyncSetUp(self):
        await test_contacts.ContactWorkflowTests.asyncSetUp(self)
        patched = patch("app.api.imports.SessionLocal", self.factory)
        patched.start()
        self.addCleanup(patched.stop)
        self.active_user = self.boss
        self.csrf = ""

    def workbook(self, rows, headers=None):
        workbook = Workbook()
        headers = headers or list(dict.fromkeys(key for row in rows for key in row))
        workbook.active.append(headers)
        for row in rows:
            workbook.active.append([row.get(key) for key in headers])
        output = BytesIO()
        workbook.save(output)
        workbook.close()
        return output.getvalue()

    async def upload(self, content, csrf=None):
        status, _, body = await self.request("GET", "/import")
        if status != 200:
            return status, {}, body
        self.csrf = html.unescape(re.search(r'name="csrf_token" value="([^"]+)"', body.decode()).group(1))
        boundary = "ImportTestBoundary"
        payload = (f"--{boundary}\r\nContent-Disposition: form-data; name=\"csrf_token\"\r\n\r\n{self.csrf if csrf is None else csrf}\r\n"
                   f"--{boundary}\r\nContent-Disposition: form-data; name=\"file\"; filename=\"clients.xlsx\"\r\n"
                   "Content-Type: application/vnd.openxmlformats-officedocument.spreadsheetml.sheet\r\n\r\n").encode() + content + f"\r\n--{boundary}--\r\n".encode()
        messages = []
        async def receive():
            return {"type": "http.request", "body": payload, "more_body": False}
        async def send(message):
            messages.append(message)
        await app({"type": "http", "asgi": {"version": "3.0", "spec_version": "2.4"}, "http_version": "1.1",
                   "method": "POST", "scheme": "http", "path": "/import", "raw_path": b"/import", "query_string": b"",
                   "root_path": "", "headers": [(b"host", b"testserver"), (b"cookie", self.cookies.encode()),
                       (b"content-type", f"multipart/form-data; boundary={boundary}".encode())],
                   "client": ("127.0.0.1", 12345), "server": ("testserver", 80)}, receive, send)
        start = next(m for m in messages if m["type"] == "http.response.start")
        headers = {k.decode(): v.decode() for k, v in start["headers"]}
        if "set-cookie" in headers:
            self.cookies = headers["set-cookie"].split(";", 1)[0]
        return start["status"], headers, b"".join(m.get("body", b"") for m in messages)

    async def preview(self, rows):
        status, headers, body = await self.upload(self.workbook(rows))
        self.assertEqual(status, 303, body.decode())
        batch_id = headers["location"].split("/")[-1]
        async with self.factory() as session:
            batch = await session.get(ClientImport, batch_id)
            return batch_id, json.loads(batch.payload)["plan"]

    async def confirm(self, batch_id):
        status, _, body = await self.request("POST", f"/import/{batch_id}/confirm", {"csrf_token": self.csrf})
        self.assertEqual(status, 303, body.decode())

    async def batch_state(self, batch_id):
        async with self.factory() as session:
            return (await session.get(ClientImport, batch_id)).state

    async def test_mixed_import_preview_apply_and_repeat(self):
        async with self.factory() as session:
            client = await session.get(Client, self.client.id)
            client.status = "Переговоры"
            client.last_call_date = self.today
            client.phone = "+79000000000"
            session.add(Comment(client_id=client.id, author_id=self.manager.id, body="Прежняя заметка", source="card"))
            session.add(ClientFunnel(client_id=client.id, funnel_id=self.funnel.id))
            await session.commit()
        rows = [
            {"ИНН": self.client.tax_number, "Организация": "Новое название", "Телефон": "+79999999999",
             "Статус": "Поиск контакта", "Дата звонка": "2026-01-01", "Следующая дата звонка": "28.10.2026",
             "Табельный номер": self.other.employee_id, "ФИО менеджера": "Игнорируем",
             "Воронки": "Спарк, Новая", "Комментарий менеджера": "Новая заметка"},
            {"ИНН": "001111111111", "Организация": "Новый клиент", "Телефон": "12345",
             "Оборот в месяц": "123,45", "Комментарий менеджера": "Первое примечание"},
            {"ИНН": "bad", "Организация": "Ошибка"},
        ]
        batch_id, plan = await self.preview(rows)
        self.assertEqual(plan["counts"], {"create": 1, "update": 1, "unchanged": 0, "error": 1, "duplicate": 0})
        self.assertEqual(len(plan["entries"][0]["warnings"]), 4)
        async with self.factory() as session:
            self.assertEqual(await session.scalar(select(func.count(Client.id))), 3)
            self.assertEqual((await session.get(Client, self.client.id)).phone, "+79000000000")
        status, _, body = await self.request("GET", f"/import/{batch_id}")
        self.assertEqual(status, 200)
        self.assertIn("Применить изменения к 2 клиентам", body.decode())
        await self.confirm(batch_id)
        await self.confirm(batch_id)
        self.assertEqual(await self.batch_state(batch_id), "applied")
        async with self.factory() as session:
            client = await session.get(Client, self.client.id)
            self.assertEqual(client.phone, "+79000000000")
            self.assertEqual(client.status, "Переговоры")
            self.assertEqual(client.last_call_date, self.today)
            self.assertEqual(client.next_call_date, date(2026, 10, 28))
            self.assertEqual(client.assigned_manager_id, self.other.id)
            self.assertIsNotNone(await session.scalar(select(Comment.id).where(Comment.body == "Прежняя заметка")))
            names = list(await session.scalars(select(Funnel.name).join(ClientFunnel).where(ClientFunnel.client_id == client.id)))
            self.assertEqual(set(names), {"Спарк", "Новая"})
            new = await session.scalar(select(Client).where(Client.tax_number == "001111111111"))
            self.assertIsNone(new.assigned_manager_id)
            self.assertEqual(new.full_name, "Новый клиент")
            self.assertEqual(str(new.monthly_turnover), "123.45")
            self.assertEqual(await session.scalar(select(func.count(Comment.id))), 3)
            self.assertEqual(await session.scalar(select(func.count(Comment.id)).where(Comment.source == "import")), 2)
            self.assertEqual(await session.scalar(select(func.count(ContactActivity.id))), 0)
        repeat_id, repeat = await self.preview(rows)
        self.assertEqual(repeat["counts"]["unchanged"], 2)
        self.assertEqual(repeat["apply_count"], 0)
        await self.confirm(repeat_id)
        async with self.factory() as session:
            self.assertEqual(await session.scalar(select(func.count(Comment.id))), 3)

    async def test_blank_cells_preserve_fields_and_equal_phone_does_not_update(self):
        async with self.factory() as session:
            client = await session.get(Client, self.client.id)
            client.phone = "+7 (999) 123-45-67"
            session.add(ClientFunnel(client_id=client.id, funnel_id=self.funnel.id))
            await session.commit()
        batch_id, plan = await self.preview([{"ИНН": self.client.tax_number, "Телефон": "89991234567",
                                             "Организация": "  ООО   Ромашка  ", "Воронки": "", "Табельный номер": "", "Статус": ""}])
        self.assertEqual(plan["counts"]["unchanged"], 1)
        await self.confirm(batch_id)
        async with self.factory() as session:
            client = await session.get(Client, self.client.id)
            self.assertEqual(client.phone, "+7 (999) 123-45-67")
            self.assertEqual(client.assigned_manager_id, self.manager.id)
            self.assertEqual(await session.scalar(select(func.count(ClientFunnel.id))), 1)

    async def test_report_has_explicit_counts_without_row_details(self):
        async with self.factory() as session:
            client = await session.get(Client, self.client.id)
            client.phone = "+79000000000"
            await session.commit()
        new = {"ИНН": "001111111111", "Организация": "Новый"}
        batch_id, plan = await self.preview([
            {"ИНН": self.client.tax_number, "Телефон": "+79999999999"},
            {"ИНН": self.second_client.tax_number, "Телефон": "12345", "Комментарий менеджера": "Новое примечание"},
            new, new, {"ИНН": "bad"}, {"ИНН": self.other_client.tax_number},
        ])
        self.assertEqual(plan["counts"], {"create": 1, "update": 1, "unchanged": 2, "error": 1, "duplicate": 1})
        status, _, body = await self.request("GET", f"/import/{batch_id}")
        self.assertEqual(status, 200)
        text = body.decode()
        self.assertNotIn("built-in", text)
        expected = {"Всего проверено строк": "6", "Новых клиентов": "1", "Клиентов с обновлениями": "1",
                    "Без изменений": "2", "Строк с ошибками": "1", "Повторных строк — пропущено": "1"}
        counters = dict(re.findall(r'<div class="stat-card"><div>(.*?)</div><strong>(\d+)</strong>', text))
        self.assertEqual(counters, expected)
        self.assertNotIn("Было в базе", text)
        self.assertNotIn('class="import-row"', text)
        self.assertNotIn("Новое примечание", text)
        self.assertNotIn("+79999999999", text)
        self.assertIn("Скачать отчёт", text)
        await self.confirm(batch_id)
        async with self.factory() as session:
            self.assertEqual((await session.get(Client, self.client.id)).phone, "+79000000000")
            self.assertEqual((await session.get(Client, self.second_client.id)).phone, "12345")
        _, _, body = await self.request("GET", f"/import/{batch_id}")
        self.assertIn("Обновлено клиентов", body.decode())
        self.assertNotIn("built-in", body.decode())

    async def test_old_preview_cannot_overwrite_phone_after_policy_change(self):
        async with self.factory() as session:
            client = await session.get(Client, self.client.id)
            client.phone = "11111"
            await session.commit()
        batch_id, _ = await self.preview([{"ИНН": self.client.tax_number, "Телефон": "22222"}])
        async with self.factory() as session:
            batch = await session.get(ClientImport, batch_id)
            data = json.loads(batch.payload)
            data["plan"].pop("version")
            data["plan"]["entries"][0]["updates"]["phone"] = "22222"
            batch.payload = json.dumps(data)
            await session.commit()
        _, _, body = await self.request("GET", f"/import/{batch_id}")
        self.assertIn("Правила импорта изменились", body.decode())
        self.assertNotIn(f'action="/import/{batch_id}/confirm"', body.decode())
        await self.confirm(batch_id)
        self.assertEqual(await self.batch_state(batch_id), "stale")
        async with self.factory() as session:
            self.assertEqual((await session.get(Client, self.client.id)).phone, "11111")

    async def test_conflicting_duplicates_reject_entire_group_identical_skip(self):
        _, plan = await self.preview([
            {"ИНН": self.client.tax_number, "Телефон": "111"},
            {"ИНН": self.client.tax_number, "Телефон": "222"},
            {"ИНН": self.second_client.tax_number, "Телефон": "333"},
            {"ИНН": self.second_client.tax_number, "Телефон": "333"},
        ])
        self.assertEqual(plan["counts"]["error"], 2)
        self.assertEqual(plan["counts"]["update"], 1)
        self.assertEqual(plan["counts"]["duplicate"], 1)

    async def test_invalid_values_are_row_errors(self):
        for changes in ({"Табельный номер": "99999999"}, {"Табельный номер": self.boss.employee_id},
                        {"Оборот в месяц": "NaN"}, {"Дата звонка": "31.02.2026"},
                        {"Статус": "Неизвестный"}, {"Организация": "=1+1"}):
            _, plan = await self.preview([{"ИНН": self.client.tax_number, **changes}])
            self.assertEqual(plan["counts"]["error"], 1, changes)

    async def test_fixed_field_policy_preserves_identity_and_updates_business_fields(self):
        async with self.factory() as session:
            client = await session.get(Client, self.client.id)
            client.phone = "11111"
            client.status = "Переговоры"
            client.last_call_date = date(2026, 10, 20)
            client.next_call_date = date(2026, 10, 30)
            client.last_manager_activity_date = date(2026, 10, 21)
            await session.commit()
        batch_id, plan = await self.preview([{
            "ИНН": self.client.tax_number, "ФИО": "Другое ФИО", "Организация": "Другая организация",
            "Телефон": "22222", "Статус": "Поиск контакта", "Дата звонка": "2026-01-01",
            "Следующая дата звонка": "2026-01-02", "Дата последней активности менеджера": "НЕ ДАТА — ИГНОРИРУЕТСЯ",
            "Оборот в месяц": "123.45", "MCC": "5411", "Банк-конкурент": "Новый банк",
            "Адрес регистрации": "Новый адрес", "Табельный номер": self.other.employee_id,
            "Воронки": "Спарк", "Комментарий менеджера": "Новая информация",
        }])
        self.assertEqual(plan["counts"]["update"], 1)
        await self.confirm(batch_id)
        async with self.factory() as session:
            client = await session.get(Client, self.client.id)
            self.assertEqual(client.full_name, self.client.full_name)
            self.assertEqual(client.company_name, self.client.company_name)
            self.assertEqual(client.tax_number, self.client.tax_number)
            self.assertEqual(client.phone, "11111")
            self.assertEqual(client.status, "Переговоры")
            self.assertEqual(client.last_call_date, date(2026, 10, 20))
            self.assertEqual(client.next_call_date, date(2026, 10, 30))
            self.assertEqual(client.last_manager_activity_date, date(2026, 10, 21))
            self.assertEqual(str(client.monthly_turnover), "123.45")
            self.assertEqual(client.mcc_code, "5411")
            self.assertEqual(client.competitor_bank, "Новый банк")
            self.assertEqual(client.registration_address, "Новый адрес")
            self.assertEqual(client.assigned_manager_id, self.other.id)
            self.assertEqual(await session.scalar(select(Comment.body).where(Comment.client_id == client.id)), "Новая информация")

    async def test_protected_empty_names_can_be_filled(self):
        async with self.factory() as session:
            client = await session.get(Client, self.client.id)
            client.company_name = ""
            client.full_name = ""
            await session.commit()
        batch_id, _ = await self.preview([{"ИНН": self.client.tax_number, "ФИО": "Новое ФИО", "Организация": "Организация"}])
        await self.confirm(batch_id)
        async with self.factory() as session:
            client = await session.get(Client, self.client.id)
            self.assertEqual(client.full_name, "Новое ФИО")
            self.assertEqual(client.company_name, "Организация")

    async def test_large_preview_page_stays_compact(self):
        batch_id, plan = await self.preview([
            {"ИНН": f"88000000{index:04d}", "Организация": f"Строка импорта {index}"}
            for index in range(3000)
        ])
        self.assertEqual(plan["counts"]["create"], 3000)
        status, _, body = await self.request("GET", f"/import/{batch_id}")
        self.assertEqual(status, 200)
        self.assertIn("3000", body.decode())
        self.assertNotIn("Строка импорта", body.decode())
        self.assertLess(len(body), 20000)
        self.assertEqual(len(await self.read_report(batch_id)), 3000)

    async def test_empty_status_dates_filled_and_funnel_names_reused(self):
        async with self.factory() as session:
            client = await session.get(Client, self.client.id)
            client.status = ""
            await session.commit()
        batch_id, plan = await self.preview([
            {"ИНН": self.client.tax_number, "Статус": "Переговоры", "Дата звонка": date(2026, 10, 25),
             "Воронки": "  спарк, Спарк  ", "Комментарий менеджера": "Заметка"},
            {"ИНН": self.second_client.tax_number, "Воронки": "НОВАЯ"},
            {"ИНН": self.other_client.tax_number, "Воронки": "новая"},
        ])
        self.assertEqual(plan["apply_count"], 3)
        await self.confirm(batch_id)
        async with self.factory() as session:
            client = await session.get(Client, self.client.id)
            self.assertEqual(client.status, "Переговоры")
            self.assertEqual(client.last_call_date, date(2026, 10, 25))
            self.assertEqual(await session.scalar(select(func.count(Funnel.id))), 3)
            self.assertEqual(await session.scalar(select(func.count(ClientFunnel.id))), 3)
        _, repeated = await self.preview([{"ИНН": self.client.tax_number, "Комментарий менеджера": "  Заметка  "}])
        self.assertEqual(repeated["counts"]["unchanged"], 1)

    async def test_database_text_is_not_executed_as_formula_on_export(self):
        async with self.factory() as session:
            client = await session.get(Client, self.client.id)
            client.company_name = "=1+1"
            await session.commit()
            content = await export_clients_to_workbook(session)
        workbook = load_workbook(BytesIO(content))
        cells = [row[1] for row in workbook.active.iter_rows(min_row=2) if row[1].value == "=1+1"]
        self.assertEqual(len(cells), 1)
        self.assertEqual(cells[0].data_type, "s")
        workbook.close()

    async def test_stale_preview_is_not_applied_even_to_unchanged_clients(self):
        batch_id, _ = await self.preview([{"ИНН": self.client.tax_number, "Телефон": "111"},
                                          {"ИНН": self.second_client.tax_number, "Телефон": "222"}])
        async with self.factory() as session:
            client = await session.get(Client, self.client.id)
            client.phone = "Changed meanwhile"
            await session.commit()
        await self.confirm(batch_id)
        self.assertEqual(await self.batch_state(batch_id), "stale")
        async with self.factory() as session:
            self.assertEqual((await session.get(Client, self.second_client.id)).phone, "")

    async def test_new_client_created_after_preview_causes_conflict(self):
        batch_id, _ = await self.preview([{"ИНН": "001111111111", "Организация": "Из файла"}])
        async with self.factory() as session:
            session.add(Client(tax_number="001111111111", company_name="Добавлен вручную", full_name="Клиент"))
            await session.commit()
        await self.confirm(batch_id)
        self.assertEqual(await self.batch_state(batch_id), "stale")

    async def test_cancel_and_expired_preview_cannot_apply(self):
        for state in ("cancelled", "stale"):
            batch_id, _ = await self.preview([{"ИНН": self.client.tax_number, "Телефон": "111"}])
            if state == "cancelled":
                status, _, _ = await self.request("POST", f"/import/{batch_id}/cancel", {"csrf_token": self.csrf})
                self.assertEqual(status, 303)
            else:
                async with self.factory() as session:
                    batch = await session.get(ClientImport, batch_id)
                    batch.created_at = datetime.utcnow() - timedelta(days=2)
                    await session.commit()
            await self.confirm(batch_id)
            self.assertEqual(await self.batch_state(batch_id), state)
        async with self.factory() as session:
            self.assertEqual((await session.get(Client, self.client.id)).phone, "")

    async def test_concurrent_confirmation_applies_once(self):
        batch_id, _ = await self.preview([{"ИНН": "001111111111", "Организация": "Новый", "Комментарий менеджера": "Один раз"}])
        await asyncio.gather(self.confirm(batch_id), self.confirm(batch_id))
        async with self.factory() as session:
            self.assertEqual(await session.scalar(select(func.count(Client.id))), 4)
            self.assertEqual(await session.scalar(select(func.count(Comment.id))), 1)

    async def test_failure_rolls_back_entire_batch(self):
        batch_id, _ = await self.preview([{"ИНН": self.client.tax_number, "Телефон": "111"}])
        async def fail(session, actor, plan):
            await apply_plan(session, actor, plan)
            await session.flush()
            raise RuntimeError("Simulated failure")
        with patch("app.api.imports.apply_plan", fail), self.assertRaises(RuntimeError):
            await self.confirm(batch_id)
        self.assertEqual(await self.batch_state(batch_id), "pending")
        async with self.factory() as session:
            self.assertEqual((await session.get(Client, self.client.id)).phone, "")

    async def test_permissions_csrf_and_batch_ownership(self):
        batch_id, _ = await self.preview([{"ИНН": self.client.tax_number}])
        status, _, _ = await self.request("POST", f"/import/{batch_id}/confirm", {"csrf_token": "bad"})
        self.assertEqual(status, 403)
        status, _, _ = await self.upload(self.workbook([{"ИНН": self.client.tax_number}]), csrf="bad")
        self.assertEqual(status, 403)
        self.active_user = self.manager
        for path in ("/import", f"/import/{batch_id}", f"/import/{batch_id}/errors", f"/import/{batch_id}/report"):
            status, _, _ = await self.request("GET", path)
            self.assertEqual(status, 403)
        self.active_user = User(id=self.other.id, role="boss")
        for method, path in (("GET", f"/import/{batch_id}"), ("GET", f"/import/{batch_id}/errors"),
                             ("GET", f"/import/{batch_id}/report"),
                             ("POST", f"/import/{batch_id}/confirm"), ("POST", f"/import/{batch_id}/cancel")):
            status, _, _ = await self.request(method, path, {"csrf_token": self.csrf} if method == "POST" else None)
            self.assertEqual(status, 403)
        app.dependency_overrides.pop(get_current_user)
        self.cookies = ""
        status, headers, _ = await self.request("GET", "/import")
        self.assertEqual(status, 303)
        self.assertEqual(headers["location"], "/login")

    async def test_report_is_safe_and_can_be_corrected_and_reuploaded(self):
        batch_id, _ = await self.preview([{"ИНН": "bad", "Организация": "=1+1", "Дата звонка": date(2026, 10, 25)}])
        status, _, body = await self.request("GET", f"/import/{batch_id}/report")
        self.assertEqual(status, 200)
        workbook = load_workbook(BytesIO(body))
        sheet = workbook.active
        self.assertEqual(sheet["B2"].data_type, "s")
        self.assertEqual(sheet["B2"].value, "=1+1")
        sheet["A2"] = "001111111111"
        sheet["B2"] = "Исправленный"
        output = BytesIO()
        workbook.save(output)
        workbook.close()
        status, headers, _ = await self.upload(output.getvalue())
        self.assertEqual(status, 303)
        async with self.factory() as session:
            plan = json.loads((await session.get(ClientImport, headers["location"].split("/")[-1])).payload)["plan"]
            self.assertEqual(plan["counts"]["create"], 1)

        # Re-exporting a corrected report must not duplicate its diagnostic columns.
        _, _, body = await self.request("GET", headers["location"] + "/report")
        workbook = load_workbook(BytesIO(body))
        titles = [cell.value for cell in workbook.active[1]]
        self.assertEqual(titles.count("Результат импорта"), 1)
        self.assertEqual(titles.count("Ошибка импорта"), 1)
        workbook.close()

    async def read_report(self, batch_id):
        status, headers, body = await self.request("GET", f"/import/{batch_id}/report")
        self.assertEqual(status, 200)
        self.assertEqual(headers["cache-control"], "no-store")
        workbook = load_workbook(BytesIO(body))
        sheet = workbook.active
        self.assertEqual(sheet.freeze_panes, "A2")
        self.assertEqual(sheet.auto_filter.ref, sheet.dimensions)
        rows = list(sheet.values)
        workbook.close()
        return [dict(zip(rows[0], row)) for row in rows[1:]]

    async def test_report_contains_all_results_changes_and_protected_fields(self):
        async with self.factory() as session:
            client = await session.get(Client, self.client.id)
            client.phone = "11111"
            client.mcc_code = "5411"
            await session.commit()
        new = {"ИНН": "001111111111", "Организация": "Новый"}
        batch_id, _ = await self.preview([
            {"ИНН": self.client.tax_number, "Телефон": "22222", "MCC": "5812",
             "Комментарий менеджера": "Новый комментарий"},
            {"ИНН": self.second_client.tax_number}, new, new, {"ИНН": "bad"},
        ])
        rows = await self.read_report(batch_id)
        self.assertEqual(len(rows), 5)
        self.assertEqual([r["Строка исходного файла"] for r in rows], [2, 3, 4, 5, 6])
        self.assertEqual([r["Результат импорта"] for r in rows],
                         ["Будет обновлён", "Без изменений", "Будет создан", "Повтор — пропущен", "Ошибка"])
        self.assertIn("MCC: 5411 → 5812", rows[0]["Изменения импорта"])
        self.assertIn("Новый комментарий", rows[0]["Изменения импорта"])
        self.assertIn("защищено", rows[0]["Пояснение импорта"])
        self.assertEqual(rows[0]["Телефон"], "22222")  # Source data, not the stored phone.
        self.assertIn("Повтор строки", rows[3]["Пояснение импорта"])
        self.assertIsNone(rows[3]["Ошибка импорта"])
        self.assertIn("ИНН", rows[4]["Ошибка импорта"])
        await self.confirm(batch_id)
        rows = await self.read_report(batch_id)
        self.assertEqual(rows[0]["Результат импорта"], "Обновлён")
        self.assertEqual(rows[2]["Результат импорта"], "Создан")
        self.assertEqual(rows[0]["Состояние загрузки"], "Импорт завершён")

    async def test_report_available_without_errors_and_labels_unapplied_batches(self):
        for state in ("cancelled", "stale", "expired", "old_policy"):
            batch_id, _ = await self.preview([{"ИНН": "001111111111", "Организация": "Новый"}])
            _, _, body = await self.request("GET", f"/import/{batch_id}")
            self.assertIn(f'href="/import/{batch_id}/report"', body.decode())
            async with self.factory() as session:
                batch = await session.get(ClientImport, batch_id)
                if state == "expired":
                    batch.created_at = datetime.utcnow() - timedelta(days=2)
                elif state == "old_policy":
                    payload = json.loads(batch.payload)
                    payload["plan"].pop("version")
                    batch.payload = json.dumps(payload)
                else:
                    batch.state = state
                await session.commit()
            rows = await self.read_report(batch_id)
            self.assertEqual(rows[0]["Результат импорта"], "Создание не применено")
            self.assertIn("изменения не применены", rows[0]["Состояние загрузки"])
            self.assertEqual(await self.batch_state(batch_id), "pending" if state in {"expired", "old_policy"} else state)

    async def test_template_route_compatibility_and_malformed_file(self):
        status, _, content = await self.request("GET", "/import/template")
        self.assertEqual(status, 200)
        workbook = load_workbook(BytesIO(content))
        self.assertNotIn("ФИО менеджера", [c.value for c in workbook.active[1]])
        self.assertNotIn("Дата последней активности менеджера", [c.value for c in workbook.active[1]])
        self.assertEqual(workbook.active["F2"].number_format, "DD.MM.YYYY")
        workbook.close()
        self.assertFalse(parse_workbook(build_import_template())["rows"][0]["error"])
        for content in (b"not excel", b"x" * (MAX_FILE_BYTES + 1), self.workbook([{"Организация": "Без ИНН"}])):
            status, _, _ = await self.upload(content)
            self.assertEqual(status, 400)
