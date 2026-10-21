from __future__ import annotations

import asyncio
import html
import json
import re
import unittest
from datetime import date, timedelta
from decimal import Decimal
from io import BytesIO
from pathlib import Path
from tempfile import TemporaryDirectory
from types import SimpleNamespace
from unittest.mock import patch
from urllib.parse import urlencode, urlsplit

from openpyxl import load_workbook
from sqlalchemy import event, func, select
from sqlalchemy.ext.asyncio import async_sessionmaker, create_async_engine

from app.api.deps import get_current_user
from app.db.session import Base, configure_sqlite
from app.main import app
from app.models import Client, ClientFunnel, Comment, ContactActivity, ContactReportState, Funnel, User
from app.services.contact_reports import REPORT_HEADERS, build_contact_workbook, retry_pending_reports
from app.services.contacts import CONTACT_RESULTS, aggregate_contacts, list_contacts


class ContactWorkflowTests(unittest.IsolatedAsyncioTestCase):
    async def asyncSetUp(self):
        directory = TemporaryDirectory()
        self.addCleanup(directory.cleanup)
        self.root = Path(directory.name)
        self.engine = create_async_engine(f"sqlite+aiosqlite:///{(self.root / 'test.sqlite3').as_posix()}")
        event.listen(self.engine.sync_engine, "connect", configure_sqlite)
        self.factory = async_sessionmaker(self.engine, expire_on_commit=False)
        async with self.engine.begin() as connection:
            await connection.run_sync(Base.metadata.create_all)
        async with self.factory() as session:
            self.manager = User(employee_id="00000001", full_name="Первый менеджер", role="manager", password_hash="test")
            self.other = User(employee_id="00000002", full_name="Второй менеджер", role="manager", password_hash="test")
            self.boss = User(employee_id="00000003", full_name="Руководитель", role="boss", password_hash="test")
            self.funnel = Funnel(name="Доля рынка")
            self.second_funnel = Funnel(name="Спарк")
            session.add_all([self.manager, self.other, self.boss, self.funnel, self.second_funnel])
            await session.flush()
            self.client = Client(full_name="Иван Иванов", company_name="ООО Ромашка", tax_number="001234567890", assigned_manager_id=self.manager.id)
            self.second_client = Client(full_name="Анна Петрова", company_name="ООО Север", tax_number="770000000002", assigned_manager_id=self.manager.id)
            self.other_client = Client(full_name="Чужой клиент", company_name="ООО Юг", tax_number="770000000003", assigned_manager_id=self.other.id)
            session.add_all([self.client, self.second_client, self.other_client])
            await session.commit()
        for target, value in [
            ("app.api.contacts.SessionLocal", self.factory),
            ("app.api.clients.SessionLocal", self.factory),
            ("app.services.contact_reports.SessionLocal", self.factory),
            ("app.services.contact_reports.settings", SimpleNamespace(reports_dir=self.root / "reports")),
        ]:
            patched = patch(target, value)
            patched.start()
            self.addCleanup(patched.stop)
        self.active_user = self.manager
        async def current_user():
            return self.active_user
        app.dependency_overrides[get_current_user] = current_user
        self.addCleanup(app.dependency_overrides.clear)
        self.cookies = ""
        self.today = date.today()

    async def asyncTearDown(self):
        await self.engine.dispose()

    async def request(self, method, url, form=None):
        """Exercise real ASGI routes/middleware without an extra HTTP dependency."""
        parsed = urlsplit(url)
        payload = urlencode(form or {}).encode()
        headers = [(b"host", b"testserver")]
        if form is not None:
            headers.append((b"content-type", b"application/x-www-form-urlencoded"))
        if self.cookies:
            headers.append((b"cookie", self.cookies.encode()))
        scope = {
            "type": "http", "asgi": {"version": "3.0", "spec_version": "2.4"}, "http_version": "1.1",
            "method": method, "scheme": "http", "path": parsed.path, "raw_path": parsed.path.encode(),
            "query_string": parsed.query.encode(), "root_path": "", "headers": headers,
            "client": ("127.0.0.1", 12345), "server": ("testserver", 80),
        }
        messages = []
        async def receive():
            return {"type": "http.request", "body": payload, "more_body": False}
        async def send(message):
            messages.append(message)
        await app(scope, receive, send)
        start = next(message for message in messages if message["type"] == "http.response.start")
        response_headers = {k.decode(): v.decode() for k, v in start["headers"]}
        if "set-cookie" in response_headers:
            self.cookies = response_headers["set-cookie"].split(";", 1)[0]
        body = b"".join(m.get("body", b"") for m in messages if m["type"] == "http.response.body")
        return start["status"], response_headers, body

    async def new_form(self, client_id=None, **changes):
        client_id = client_id or self.client.id
        status, _, body = await self.request("GET", f"/clients/{client_id}/contacts/new")
        self.assertEqual(status, 200)
        text = body.decode()
        self.assertNotIn('name="dialogue"', text)
        fields = {
            "activity_date": self.today.isoformat(), "activity_type": "Звонок",
            "result": "Переговоры", "potential_turnover": "000000", "funnel_id": str(self.funnel.id),
            "next_activity_date": "", "comment": "Обсудили условия",
        }
        for name in ("csrf_token", "submission_token"):
            fields[name] = html.unescape(re.search(rf'name="{name}" value="([^"]+)"', text).group(1))
        fields.update(changes)
        return fields

    async def save(self, client_id=None, **changes):
        client_id = client_id or self.client.id
        fields = await self.new_form(client_id, **changes)
        status, headers, body = await self.request("POST", f"/clients/{client_id}/contacts", fields)
        self.assertEqual(status, 303, body.decode())
        return headers

    async def report_rows(self, manager_id=None, start=None, end=None):
        async with self.factory() as session:
            return aggregate_contacts(await list_contacts(session, manager_id or self.manager.id, start or self.today, end or self.today))

    async def unscheduled_section(self, day=None):
        status, _, body = await self.request("GET", f"/contacts?day={(day or self.today).isoformat()}")
        self.assertEqual(status, 200)
        section = re.search(r'<section class="panel" id="unscheduled-clients">(.*?)</section>', body.decode(), re.S)
        self.assertIsNotNone(section)
        return section.group(1)

    async def test_export_preview_counts_visible_clients(self):
        with patch("app.main.SessionLocal", self.factory), patch("app.main.export_clients_to_workbook") as build:
            for user, expected in ((self.manager, 2), (self.other, 1), (self.boss, 3)):
                self.active_user = user
                status, headers, body = await self.request("GET", "/export?preview=true")
                self.assertEqual(status, 200)
                self.assertEqual(json.loads(body), {"count": expected})
                self.assertEqual(headers["cache-control"], "no-store")
            build.assert_not_called()

    async def test_navigation_hides_personal_contacts_for_boss(self):
        for user in (self.manager, self.boss):
            self.active_user = user
            status, _, body = await self.request("GET", "/contacts/journal")
            self.assertEqual(status, 200)
            navigation = re.search(r'<nav class="nav">(.*?)</nav>', body.decode(), re.S).group(1)
            if user.role == "boss":
                self.assertNotIn('href="/contacts"', navigation)
            else:
                self.assertIn('href="/contacts">Мои звонки и встречи</a>', navigation)
            self.assertIn('href="/contacts/journal"', navigation)

    async def test_boss_has_no_personal_contact_controls_or_routes(self):
        fields = await self.new_form()
        self.active_user = self.boss
        with patch("app.main.SessionLocal", self.factory):
            for path in ("/clients/new", f"/clients/{self.client.id}", "/contacts/journal"):
                status, _, body = await self.request("GET", path)
                self.assertEqual(status, 200)
                text = body.decode()
                self.assertNotIn("Мои звонки и встречи", text)
                self.assertNotIn('href="/contacts"', text)
                self.assertNotIn(f'/clients/{self.client.id}/contacts/new', text)
            for method, path, form in (
                ("GET", "/contacts", None),
                ("GET", f"/clients/{self.client.id}/contacts/new", None),
                ("POST", f"/clients/{self.client.id}/contacts", fields),
                ("POST", "/contacts/report/retry", {"day": self.today.isoformat(), "csrf_token": fields["csrf_token"]}),
            ):
                status, _, _ = await self.request(method, path, form)
                self.assertEqual(status, 403, path)
        async with self.factory() as session:
            for model in (ContactActivity, ContactReportState):
                self.assertEqual(await session.scalar(select(func.count()).select_from(model)), 0)

    async def test_service_rejects_contact_created_by_boss(self):
        from app.services.contacts import parse_contact, record_contact
        data = parse_contact(await self.new_form())
        async with self.factory() as session:
            client = await session.get(Client, self.client.id)
            with self.assertRaises(PermissionError):
                await record_contact(session, client, self.boss, data, "boss-forbidden")
            self.assertEqual(await session.scalar(select(func.count()).select_from(ContactActivity)), 0)

    async def test_boss_report_selector_excludes_non_managers_without_history(self):
        await self.save()
        # Promoting a manager must not hide the contacts they previously recorded.
        async with self.factory() as session:
            manager = await session.get(User, self.manager.id)
            manager.role = "boss"
            await session.commit()
        self.active_user = self.boss
        status, _, body = await self.request("GET", "/contacts/journal")
        self.assertEqual(status, 200)
        selector = re.search(r'<select name="manager_id".*?</select>', body.decode(), re.S).group(0)
        self.assertIn(f'value="{self.manager.id}"', selector)
        self.assertIn(f'value="{self.other.id}"', selector)
        self.assertNotIn(f'value="{self.boss.id}"', selector)
        status, _, body = await self.request("GET", f"/contacts/export?preview=true&manager_id={self.manager.id}")
        self.assertEqual(status, 200)
        self.assertEqual(json.loads(body), {"count": 1})
        for path in ("/contacts/journal", "/contacts/export"):
            status, _, _ = await self.request("GET", f"{path}?manager_id={self.boss.id}")
            self.assertEqual(status, 403)

    async def test_main_export_requires_confirmation_without_javascript(self):
        with patch("app.main.SessionLocal", self.factory):
            with patch("app.main.export_clients_to_workbook") as build:
                for user, expected in ((self.manager, 2), (self.boss, 3)):
                    self.active_user = user
                    status, headers, body = await self.request("GET", "/export")
                    self.assertEqual(status, 200)
                    self.assertIn("text/html", headers["content-type"])
                    self.assertNotIn("content-disposition", headers)
                    self.assertIn(f"Выгрузить {expected} записей клиентов?", body.decode())
                    self.assertIn('href="/export?confirmed=true"', body.decode())
                    self.assertIn('href="/dashboard">Отмена</a>', body.decode())
                    self.assertIn("app.js?v=export-confirm-2", body.decode())
                build.assert_not_called()
            self.active_user = self.manager
            status, headers, body = await self.request("GET", "/export?confirmed=true")
            self.assertEqual(status, 200)
            self.assertIn("attachment", headers["content-disposition"])
            workbook = load_workbook(BytesIO(body))
            self.assertEqual(workbook.active.max_row, 3)
            workbook.close()

    async def test_contact_export_preview_counts_aggregated_rows_and_checks_access(self):
        status, _, body = await self.request("GET", "/contacts/export?preview=true")
        self.assertEqual(status, 200)
        self.assertEqual(json.loads(body), {"count": 0})
        await self.save(comment="Первый звонок")
        await self.save(comment="Повторный звонок")
        await self.save(self.second_client.id, result="Недозвон", next_activity_date=self.today.isoformat())
        earlier = self.today.replace(day=1) - timedelta(days=1)
        await self.save(activity_date=earlier.isoformat())
        with patch("app.api.contacts.build_contact_workbook") as build:
            for period in (f"day={self.today.isoformat()}", f"month={self.today:%Y-%m}"):
                status, _, body = await self.request("GET", f"/contacts/export?preview=true&{period}")
                self.assertEqual(status, 200)
                self.assertEqual(json.loads(body), {"count": 2})
            status, _, _ = await self.request("GET", f"/contacts/export?preview=true&manager_id={self.other.id}")
            self.assertEqual(status, 403)
            self.active_user = self.boss
            status, _, body = await self.request("GET", f"/contacts/export?preview=true&manager_id={self.manager.id}&month={earlier:%Y-%m}")
            self.assertEqual(status, 200)
            self.assertEqual(json.loads(body), {"count": 1})
            build.assert_not_called()

    async def test_unscheduled_list_contains_only_owned_clients_without_date(self):
        section = await self.unscheduled_section()
        self.assertIn(self.client.tax_number, section)
        self.assertIn(self.second_client.tax_number, section)
        self.assertNotIn(self.other_client.tax_number, section)
        self.assertIn(f'href="/clients/{self.client.id}"', section)
        self.assertIn("Без следующей даты · 2", section)
        for scheduled_day in (self.today - timedelta(days=1), self.today, self.today + timedelta(days=1)):
            async with self.factory() as session:
                client = await session.get(Client, self.second_client.id)
                client.next_call_date = scheduled_day
                await session.commit()
            section = await self.unscheduled_section(self.today - timedelta(days=10))
            self.assertIn(self.client.tax_number, section)
            self.assertNotIn(self.second_client.tax_number, section)
            self.assertIn("Без следующей даты · 1", section)

    async def test_contact_without_next_date_remains_visible_and_card_can_schedule_it(self):
        await self.save(next_activity_date="")
        section = await self.unscheduled_section()
        self.assertIn(self.client.tax_number, section)
        with patch("app.main.SessionLocal", self.factory):
            status, _, body = await self.request("GET", f"/clients/{self.client.id}")
            self.assertEqual(status, 200)
            self.assertIn('name="next_call_date"', body.decode())
            status, _, _ = await self.request("POST", f"/clients/{self.client.id}/save", {
                "next_call_date": self.today.isoformat(),
            })
            self.assertEqual(status, 303)
            self.assertNotIn(self.client.tax_number, await self.unscheduled_section())
            _, _, body = await self.request("GET", "/contacts")
            plan = re.search(r'<section class="panel"><h2>Назначено на .*?</section>', body.decode(), re.S).group(0)
            self.assertIn(self.client.tax_number, plan)
            status, _, _ = await self.request("POST", f"/clients/{self.client.id}/save", {"next_call_date": ""})
            self.assertEqual(status, 303)
        self.assertIn(self.client.tax_number, await self.unscheduled_section())
        rows = await self.report_rows()
        self.assertEqual(sum(row.qualified for row in rows), 1)

    async def test_unscheduled_list_empty_state(self):
        async with self.factory() as session:
            for client_id in (self.client.id, self.second_client.id):
                client = await session.get(Client, client_id)
                client.next_call_date = self.today
            await session.commit()
        section = await self.unscheduled_section()
        self.assertIn("Клиентов без следующей даты нет.", section)
        self.assertNotIn('class="unscheduled-client-grid unscheduled-client-row"', section)

    async def new_client_fields(self, **changes):
        status, _, body = await self.request("GET", f"/clients/new?funnel_id={self.funnel.id}")
        self.assertEqual(status, 200)
        text = body.decode()
        self.assertIn(f'value="{self.funnel.id}" selected', text)
        fields = {
            "tax_number": "001111111111", "company_name": "Новая организация",
            "full_name": "Контактное лицо", "phone": "+7 900 123-45-67",
            "funnel_id": str(self.funnel.id),
            "csrf_token": html.unescape(re.search(r'name="csrf_token" value="([^"]+)"', text).group(1)),
        }
        fields.update(changes)
        return fields

    async def test_manager_creates_client_in_funnel_and_cannot_override_owner(self):
        fields = await self.new_client_fields(assigned_manager_id=str(self.other.id), status="Подписан")
        status, headers, _ = await self.request("POST", "/clients/new", fields)
        self.assertEqual(status, 303)
        async with self.factory() as session:
            client = (await session.execute(select(Client).where(Client.tax_number == fields["tax_number"]))).scalar_one()
            self.assertEqual(client.assigned_manager_id, self.manager.id)
            self.assertEqual(client.company_name, fields["company_name"])
            self.assertEqual(client.full_name, fields["full_name"])
            self.assertEqual(client.phone, fields["phone"])
            self.assertEqual(client.status, "Поиск контакта")
            links = list((await session.execute(select(ClientFunnel).where(ClientFunnel.client_id == client.id))).scalars())
            self.assertEqual([link.funnel_id for link in links], [self.funnel.id])
            self.assertEqual(headers["location"], f"/clients/{client.id}?funnel_id={self.funnel.id}&created=1")
            client_id = client.id
            from app.services.task_service import get_funnel_clients
            clients = await get_funnel_clients(self.funnel.id, self.manager, session)
            self.assertIn(client_id, [item.id for item in clients])
        self.assertEqual(await self.report_rows(), [])
        await self.save(client_id, comment="Первый звонок новому клиенту")
        self.assertTrue((await self.report_rows())[0].qualified)

    async def test_boss_creates_unassigned_client_and_card_preserves_empty_assignment(self):
        self.active_user = self.boss
        fields = await self.new_client_fields(assigned_manager_id=str(self.manager.id))
        status, headers, _ = await self.request("POST", "/clients/new", fields)
        self.assertEqual(status, 303)
        async with self.factory() as session:
            client = (await session.execute(select(Client).where(Client.tax_number == fields["tax_number"]))).scalar_one()
            client_id = client.id
            self.assertIsNone(client.assigned_manager_id)
            links = list((await session.execute(select(ClientFunnel).where(ClientFunnel.client_id == client_id))).scalars())
            self.assertEqual([link.funnel_id for link in links], [self.funnel.id])
        with patch("app.main.SessionLocal", self.factory):
            status, _, body = await self.request("GET", headers["location"])
            self.assertEqual(status, 200)
            self.assertIn('<option value="" selected>Не назначен</option>', body.decode())
            self.assertIn("без закрепления за менеджером", body.decode())
            status, _, _ = await self.request("POST", f"/clients/{client_id}/save", {
                "assigned_manager_id": "", "phone": "12345",
            })
            self.assertEqual(status, 303)
        async with self.factory() as session:
            client = await session.get(Client, client_id)
            self.assertIsNone(client.assigned_manager_id)
            self.assertEqual(client.phone, "12345")

    async def test_boss_can_unassign_and_reassign_without_changing_on_missing_field(self):
        self.active_user = self.boss
        with patch("app.main.SessionLocal", self.factory):
            for fields, expected in (
                ({"phone": "12345"}, self.manager.id),
                ({"assigned_manager_id": ""}, None),
                ({"assigned_manager_id": str(self.other.id)}, self.other.id),
                ({"phone": "67890"}, self.other.id),
            ):
                status, _, _ = await self.request("POST", f"/clients/{self.client.id}/save", fields)
                self.assertEqual(status, 303)
                async with self.factory() as session:
                    client = await session.get(Client, self.client.id)
                    self.assertEqual(client.assigned_manager_id, expected)

    async def test_manager_cannot_clear_or_change_assignment(self):
        with patch("app.main.SessionLocal", self.factory):
            for value in ("", str(self.other.id)):
                status, _, _ = await self.request("POST", f"/clients/{self.client.id}/save", {"assigned_manager_id": value})
                self.assertEqual(status, 303)
                async with self.factory() as session:
                    client = await session.get(Client, self.client.id)
                    self.assertEqual(client.assigned_manager_id, self.manager.id)

    async def test_legacy_boss_assignment_is_not_silently_replaced_in_card(self):
        async with self.factory() as session:
            client = await session.get(Client, self.client.id)
            client.assigned_manager_id = self.boss.id
            await session.commit()
        self.active_user = self.boss
        with patch("app.main.SessionLocal", self.factory):
            status, _, body = await self.request("GET", f"/clients/{self.client.id}")
            self.assertEqual(status, 200)
            self.assertIn(f'<option value="{self.boss.id}" selected disabled>', body.decode())
            status, _, _ = await self.request("POST", f"/clients/{self.client.id}/save", {"phone": "12345"})
            self.assertEqual(status, 303)
        async with self.factory() as session:
            client = await session.get(Client, self.client.id)
            self.assertEqual(client.assigned_manager_id, self.boss.id)

    async def test_duplicate_inn_does_not_reassign_or_change_client(self):
        async with self.factory() as session:
            unassigned = Client(full_name="Без менеджера", company_name="Свободный клиент", tax_number="001111111112")
            session.add(unassigned)
            await session.commit()
        for target in (self.client, self.other_client, unassigned):
            with self.subTest(client_id=target.id):
                fields = await self.new_client_fields(tax_number=" " + target.tax_number + " ", company_name="Не перезаписывать")
                status, _, body = await self.request("POST", "/clients/new", fields)
                self.assertEqual(status, 409)
                if target.assigned_manager_id == self.manager.id:
                    self.assertIn(f'href="/clients/{target.id}"', body.decode())
                else:
                    self.assertNotIn(f'href="/clients/{target.id}"', body.decode())
                    self.assertNotIn(target.company_name, body.decode())
                async with self.factory() as session:
                    client = await session.get(Client, target.id)
                    self.assertEqual(client.assigned_manager_id, target.assigned_manager_id)
                    self.assertEqual(client.company_name, target.company_name)
                    self.assertEqual((await session.execute(select(func.count(ClientFunnel.id)))).scalar(), 0)

    async def test_concurrent_creation_of_same_inn_creates_one_client(self):
        fields = await self.new_client_fields()
        responses = await asyncio.gather(
            self.request("POST", "/clients/new", fields),
            self.request("POST", "/clients/new", fields),
        )
        self.assertEqual(sorted(response[0] for response in responses), [303, 409])
        async with self.factory() as session:
            self.assertEqual((await session.execute(select(func.count(Client.id)).where(Client.tax_number == fields["tax_number"]))).scalar(), 1)
            self.assertEqual((await session.execute(select(func.count(ClientFunnel.id)))).scalar(), 1)

    async def test_new_client_validation_and_optional_fields(self):
        for changes in (
            {"tax_number": ""}, {"tax_number": "123"}, {"tax_number": "12345678901"},
            {"tax_number": "１２３４５６７８９０"}, {"tax_number": "123456789a"},
            {"company_name": " "}, {"funnel_id": ""}, {"funnel_id": "999999"},
            {"full_name": "а" * 256}, {"phone": "1" * 65},
        ):
            with self.subTest(changes=changes):
                fields = await self.new_client_fields(**changes)
                status, _, body = await self.request("POST", "/clients/new", fields)
                self.assertEqual(status, 400)
                self.assertIn('name="tax_number"', body.decode())
        async with self.factory() as session:
            self.assertEqual((await session.execute(select(func.count(Client.id)))).scalar(), 3)
            self.assertEqual((await session.execute(select(func.count(ClientFunnel.id)))).scalar(), 0)
        fields = await self.new_client_fields(full_name="", phone="", tax_number="0011111111")
        status, _, _ = await self.request("POST", "/clients/new", fields)
        self.assertEqual(status, 303)
        async with self.factory() as session:
            client = (await session.execute(select(Client).where(Client.tax_number == fields["tax_number"]))).scalar_one()
            self.assertEqual(client.full_name, fields["company_name"])

    async def test_new_client_requires_csrf_and_authentication(self):
        fields = await self.new_client_fields(csrf_token="wrong")
        status, _, _ = await self.request("POST", "/clients/new", fields)
        self.assertEqual(status, 403)
        dependency = app.dependency_overrides.pop(get_current_user)
        try:
            for method in ("GET", "POST"):
                status, headers, _ = await self.request(method, "/clients/new", fields if method == "POST" else None)
                self.assertEqual(status, 303)
                self.assertEqual(headers["location"], "/login")
        finally:
            app.dependency_overrides[get_current_user] = dependency

    async def test_full_day_merge_meeting_priority_and_daily_file(self):
        tomorrow = (self.today + timedelta(days=1)).isoformat()
        await self.save(result="Недозвон", next_activity_date=tomorrow, comment="Не ответил")
        self.assertFalse((await self.report_rows())[0].qualified)
        await self.save(result="Перенос на дату", next_activity_date=tomorrow, comment="Обсудили условия", potential_turnover="300000")
        await self.save(activity_type="Встреча", result="Договор подписан", comment="Подписали договор", potential_turnover="500000", funnel_id=str(self.second_funnel.id))
        await self.save(result="Недозвон", next_activity_date=tomorrow, comment="Не ответил повторно", potential_turnover="450000")
        rows = await self.report_rows()
        self.assertEqual(len(rows), 1)
        row = rows[0]
        self.assertTrue(row.qualified)
        self.assertEqual(row.kind, "Встреча")
        self.assertEqual(row.status, "Недозвон")
        self.assertEqual(row.potential, Decimal("450000"))
        self.assertEqual(row.channel, self.funnel.name)
        self.assertEqual(row.next_day, self.today + timedelta(days=1))
        self.assertEqual(row.comment, "Не ответил = Обсудили условия = Подписали договор = Не ответил повторно")
        paths = list((self.root / "reports").rglob("*.xlsx"))
        self.assertEqual(len(paths), 1)
        workbook = load_workbook(paths[0])
        self.addCleanup(workbook.close)
        sheet = workbook.active
        self.assertEqual(sheet.max_column, 9)
        self.assertEqual([c.value for c in sheet[3]], list(REPORT_HEADERS))
        self.assertEqual(sheet.max_row, 4)
        self.assertEqual(sheet["A4"].value, "001234567890")
        self.assertEqual(sheet["A4"].data_type, "s")
        self.assertEqual(sheet["C4"].value, 450000)
        self.assertEqual(sheet["F4"].value.date(), self.today)
        self.assertEqual(sheet["I4"].value, row.comment)
        self.assertEqual(sheet.freeze_panes, "A4")
        status, _, page = await self.request("GET", "/contacts")
        self.assertEqual(status, 200)
        self.assertIn("1 из 20", page.decode())
        async with self.factory() as session:
            client = await session.get(Client, self.client.id)
            self.assertEqual(client.status, "Подписан")
            self.assertEqual(client.next_call_date, self.today + timedelta(days=1))

    async def test_unique_clients_and_monthly_grouping(self):
        await self.save()
        await self.save(comment="Повторный разговор")
        await self.save(self.second_client.id, activity_type="Встреча", comment="Встреча")
        yesterday = self.today - timedelta(days=1)
        await self.save(activity_date=yesterday.isoformat(), comment="Вчерашний разговор")
        rows = await self.report_rows(start=yesterday)
        self.assertEqual(len(rows), 3)
        self.assertEqual(sum(row.qualified for row in rows if row.day == self.today), 2)
        daily_files = list((self.root / "reports").rglob("*.xlsx"))
        self.assertEqual(len(daily_files), 2)
        for path in daily_files:
            workbook = load_workbook(path)
            days = {workbook.active.cell(i, 6).value.date() for i in range(4, workbook.active.max_row + 1)}
            self.assertEqual(len(days), 1)
            workbook.close()
        status, _, _ = await self.request("GET", f"/contacts/journal?month={self.today:%Y-%m}")
        self.assertEqual(status, 200)

    async def test_idempotency_and_history_survives_reassignment(self):
        fields = await self.new_form()
        for _ in range(2):
            status, _, _ = await self.request("POST", f"/clients/{self.client.id}/contacts", fields)
            self.assertEqual(status, 303)
        async with self.factory() as session:
            self.assertEqual((await session.execute(select(func.count(ContactActivity.id)))).scalar(), 1)
            self.assertEqual((await session.execute(select(func.count(Comment.id)))).scalar(), 1)
            client = await session.get(Client, self.client.id)
            client.assigned_manager_id = self.other.id
            client.company_name = "Новое имя"
            client.tax_number = "999999999999"
            funnel = await session.get(Funnel, self.funnel.id)
            funnel.name = "Переименованная воронка"
            await session.commit()
        row = (await self.report_rows())[0]
        self.assertEqual(row.company_name, "ООО Ромашка")
        self.assertEqual(row.tax_number, "001234567890")
        self.assertEqual(row.channel, "Доля рынка")
        self.assertEqual(await self.report_rows(manager_id=self.other.id), [])

    async def test_simultaneous_submission_of_same_form_records_once(self):
        fields = await self.new_form()
        responses = await asyncio.gather(
            self.request("POST", f"/clients/{self.client.id}/contacts", fields),
            self.request("POST", f"/clients/{self.client.id}/contacts", fields),
        )
        self.assertEqual([response[0] for response in responses], [303, 303])
        async with self.factory() as session:
            self.assertEqual((await session.execute(select(func.count(ContactActivity.id)))).scalar(), 1)

    async def test_unauthenticated_requests_redirect_to_login(self):
        dependency = app.dependency_overrides.pop(get_current_user)
        try:
            for path in ("/contacts", "/contacts/journal", "/contacts/export", f"/clients/{self.client.id}/contacts/new"):
                status, headers, _ = await self.request("GET", path)
                self.assertEqual(status, 303)
                self.assertEqual(headers["location"], "/login")
        finally:
            app.dependency_overrides[get_current_user] = dependency

    async def test_reporting_period_limits_exclude_other_months(self):
        this_month = self.today.replace(day=1)
        previous_month = this_month - timedelta(days=1)
        await self.save(comment="Текущий месяц")
        await self.save(activity_date=previous_month.isoformat(), comment="Предыдущий месяц")
        status, _, body = await self.request("GET", f"/contacts/export?month={self.today:%Y-%m}")
        self.assertEqual(status, 200)
        workbook = load_workbook(BytesIO(body))
        self.assertEqual(workbook.active.max_row, 4)
        self.assertEqual(workbook.active["I4"].value, "Текущий месяц")
        workbook.close()
        for path in ("/contacts?day=invalid", "/contacts/journal?month=2026-13", "/contacts/export?month=invalid"):
            status, _, _ = await self.request("GET", path)
            self.assertEqual(status, 400)

    async def test_permissions_csrf_and_boss_selection(self):
        status, _, _ = await self.request("GET", f"/clients/{self.other_client.id}/contacts/new")
        self.assertEqual(status, 403)
        fields = await self.new_form()
        status, _, _ = await self.request("POST", f"/clients/{self.other_client.id}/contacts", fields)
        self.assertEqual(status, 400)  # signed form cannot be moved to another client
        fields["csrf_token"] = "wrong"
        status, _, _ = await self.request("POST", f"/clients/{self.client.id}/contacts", fields)
        self.assertEqual(status, 403)
        for path in ("/contacts/export", "/contacts/journal"):
            status, _, _ = await self.request("GET", f"{path}?manager_id={self.other.id}&month={self.today:%Y-%m}")
            self.assertEqual(status, 403)
        self.active_user = self.boss
        status, _, body = await self.request("GET", "/contacts/journal")
        self.assertEqual(status, 200)
        self.assertIn("Выберите менеджера", body.decode())
        status, _, _ = await self.request("GET", "/contacts/export")
        self.assertEqual(status, 400)
        status, _, body = await self.request("GET", f"/contacts/export?manager_id={self.manager.id}&month={self.today:%Y-%m}")
        self.assertEqual(status, 200)
        workbook = load_workbook(BytesIO(body))
        self.assertEqual(workbook.active.max_column, 9)
        workbook.close()

    async def test_validation_preserves_form_and_creates_no_records(self):
        for changes in (
            {"result": "Недозвон"},
            {"result": "Перенос на дату"},
            {"comment": "  "}, {"potential_turnover": "NaN"}, {"potential_turnover": "-1"},
            {"potential_turnover": "0.001"}, {"potential_turnover": "1000000000000"},
            {"funnel_id": "999999"},
            {"activity_type": "Встреча", "result": "Недозвон"},
            {"activity_date": (self.today + timedelta(days=1)).isoformat()},
            {"next_activity_date": (self.today - timedelta(days=1)).isoformat()},
        ):
            with self.subTest(changes=changes):
                fields = await self.new_form(**changes)
                status, _, body = await self.request("POST", f"/clients/{self.client.id}/contacts", fields)
                self.assertEqual(status, 400)
                self.assertIn('name="submission_token"', body.decode())
        self.assertEqual(await self.report_rows(), [])

    async def test_every_status_except_no_answer_counts_without_dialogue_field(self):
        for index, result in enumerate(CONTACT_RESULTS):
            with self.subTest(result=result):
                day = self.today - timedelta(days=index)
                await self.save(result=result, activity_date=day.isoformat(), next_activity_date=self.today.isoformat())
                row = (await self.report_rows(start=day, end=day))[0]
                self.assertEqual(row.qualified, result != "Недозвон")
                async with self.factory() as session:
                    contact = (await list_contacts(session, self.manager.id, day, day))[0]
                    self.assertEqual(contact.dialogue, result != "Недозвон")

    async def test_status_overrides_legacy_dialogue_flags(self):
        await self.save(result="Недозвон", dialogue="yes", next_activity_date=self.today.isoformat())
        async with self.factory() as session:
            contact = (await list_contacts(session, self.manager.id, self.today, self.today))[0]
            contact.dialogue = True
            await session.commit()
        self.assertFalse((await self.report_rows())[0].qualified)
        await self.save(result="Отказ", dialogue="no")
        async with self.factory() as session:
            contacts = await list_contacts(session, self.manager.id, self.today, self.today)
            for contact in contacts:
                contact.dialogue = False
            await session.commit()
        self.assertTrue((await self.report_rows())[0].qualified)

    async def test_failed_file_is_retried_and_contact_is_not_lost(self):
        await self.save(comment="Первая запись")
        destination = next((self.root / "reports").rglob("*.xlsx"))
        before = destination.read_bytes()
        with patch("app.services.contact_reports.os.replace", side_effect=PermissionError("Excel is open")):
            with self.assertLogs("app.services.contact_reports", level="ERROR"):
                headers = await self.save(comment="Вторая запись")
        self.assertIn("report=pending", headers["location"])
        self.assertEqual(destination.read_bytes(), before)
        self.assertEqual(list(destination.parent.glob("*.tmp")), [])
        self.assertIn("Вторая запись", (await self.report_rows())[0].comment)
        async with self.factory() as session:
            self.assertTrue((await session.get(ContactReportState, (self.manager.id, self.today))).pending)
        await retry_pending_reports()
        async with self.factory() as session:
            self.assertFalse((await session.get(ContactReportState, (self.manager.id, self.today))).pending)
        workbook = load_workbook(destination)
        self.assertEqual(workbook.active["I4"].value, "Первая запись = Вторая запись")
        workbook.close()

    async def test_historical_contact_does_not_replace_current_plan(self):
        future = self.today + timedelta(days=5)
        await self.save(result="Перенос на дату", next_activity_date=future.isoformat())
        await self.save(activity_date=(self.today - timedelta(days=1)).isoformat(), result="Отказ")
        async with self.factory() as session:
            client = await session.get(Client, self.client.id)
            self.assertEqual(client.next_call_date, future)
            self.assertEqual(client.last_call_date, self.today)
        status, _, body = await self.request("GET", f"/contacts?day={future.isoformat()}")
        self.assertEqual(status, 200)
        self.assertIn("ООО Ромашка", body.decode())
        status, _, body = await self.request("GET", f"/contacts?day={(future + timedelta(days=1)).isoformat()}")
        self.assertEqual(status, 200)
        self.assertIn("ООО Ромашка", body.decode())

    async def test_workbook_text_is_not_executed_as_formula(self):
        await self.save(comment="=HYPERLINK(\"https://example.com\")", potential_turnover="000000")
        rows = await self.report_rows()
        workbook = load_workbook(BytesIO(build_contact_workbook(rows, "Менеджер", "День")))
        self.assertEqual(workbook.active["I4"].data_type, "s")
        self.assertEqual(workbook.active["C4"].value, 0)
        self.assertIn("000000", workbook.active["C4"].number_format)
        workbook.close()

    async def test_initialization_adds_tables_to_existing_database(self):
        from app.db.session import init_db
        async with self.engine.begin() as connection:
            await connection.run_sync(ContactReportState.__table__.drop)
            await connection.run_sync(ContactActivity.__table__.drop)
        with patch("app.db.session.engine", self.engine):
            await init_db()
            await init_db()
        async with self.factory() as session:
            self.assertIsNotNone(await session.get(Client, self.client.id))
            self.assertEqual((await session.execute(select(func.count(ContactActivity.id)))).scalar(), 0)


if __name__ == "__main__":
    unittest.main()
